"""Builds Investor_Analyst_Report.docx in JDF 2.2 format. Numbers come from model.py."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from model import run, co, hess_adj

o = run()
X, C = o["XOM"], o["CVX"]
xr = {r["B"]: r for r in X["rows"]}
cr = {r["B"]: r for r in C["rows"]}
FONT = "Palatino Linotype"

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin, s.bottom_margin = Inches(1), Inches(1.5)
    s.left_margin = s.right_margin = Inches(1.5)
    s.footer_distance = Inches(1)

st = doc.styles["Normal"]
st.font.name = FONT
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.font.size = Pt(11)
pf = st.paragraph_format
pf.line_spacing = 1.26
pf.space_after = Pt(8.5)
pf.space_before = Pt(0)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

cp = doc.core_properties
cp.author = ""; cp.last_modified_by = ""; cp.title = ""; cp.subject = ""; cp.keywords = ""; cp.comments = ""

# page number footer
fp = doc.sections[0].footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; r._r.append(e)


def runs(p, text, size=None):
    """Tiny markup: **bold**, *italic*."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            rr = p.add_run(part[2:-2]); rr.bold = True
        elif part.startswith("*"):
            rr = p.add_run(part[1:-1]); rr.italic = True
        else:
            rr = p.add_run(part)
        if size:
            rr.font.size = Pt(size)
    return p


def para(text, **kw):
    return runs(doc.add_paragraph(), text, **kw)


def h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(11); p.paragraph_format.space_after = Pt(8.5)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.keep_with_next = True
    rr = p.add_run(text.upper()); rr.bold = True


def h2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8.5); p.paragraph_format.space_after = Pt(8.5)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.keep_with_next = True
    rr = p.add_run(text); rr.bold = True


def caption(label, text, keep_next=False):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.left_indent = f.right_indent = Inches(1)
    f.line_spacing_rule = WD_LINE_SPACING.EXACTLY; f.line_spacing = Pt(14)
    f.space_after = Pt(8.5); f.keep_with_next = keep_next
    rr = p.add_run(label + "—"); rr.bold = True; rr.italic = True; rr.font.size = Pt(8.5)
    runs(p, text, size=8.5)


def figure(path, label, text, width=5.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width))
    caption(label, text)


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge, val in edges.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:color"), "000000")
        b.append(e)
    tcPr.append(b)


def table(label, text, header, rows, widths):
    caption(label, text, keep_next=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, wdt in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(wdt * 1440)))
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = (i < len(rows))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            rr = p.add_run(val); rr.font.size = Pt(8.5); rr.bold = (i == 0)
            if i == 0:
                set_cell_border(cell, top=8, bottom=4)
            elif i == len(rows):
                set_cell_border(cell, bottom=8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def f1(v): return f"{v:.1f}"

# ---------------------------------------------------------------- title block
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing = 1.15
rr = p.add_run("Chevron Over ExxonMobil as the War Premium Fades"); rr.font.size = Pt(17)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Investor / Analyst Report · Energy: U.S. Integrated Oil\nData and analysis current through October 13, 2026")

ab = doc.add_paragraph()
ab.paragraph_format.left_indent = ab.paragraph_format.right_indent = Inches(0.5)
rr = ab.add_run("Abstract—"); rr.bold = True; rr.italic = True
runs(ab, (
    "Brent crude is near $105 because of the Iran conflict, and the EIA expects most of that premium to fade by "
    "late 2027. We ask whether ExxonMobil (XOM) or Chevron (CVX) is the better holding for an investor who wants "
    "integrated oil exposure without betting on the war lasting. After normalizing cash flow to a no-war "
    f"baseline, Chevron has the higher free cash flow (FCF) yield at $74 Brent ({cr[74]['fcf_yld']:.1f}% versus "
    f"{xr[74]['fcf_yld']:.1f}%) and in three of four scenarios. Exxon is the steadier stock if oil falls hard. "
    "**We prefer CVX**, unless oil settles below about $67."))

h1("1 Question and recommendation")
para("**Question.** For an investor adding a U.S. integrated oil major today, is ExxonMobil or Chevron the "
     "better choice when current oil prices carry a war premium that is expected to shrink?")
para(f"**Recommendation: prefer Chevron.** Once the war windfall is stripped out, Chevron produces more free "
     f"cash flow per dollar of market value ({cr[74]['fcf_yld']:.1f}% versus {xr[74]['fcf_yld']:.1f}% at $74) "
     "and pays the higher dividend (3.3% versus 2.4%). That holds even though our model gives Chevron no credit "
     "for its Tengiz expansion. The cost is risk. Chevron's cash flow is more sensitive to oil, and Exxon has a "
     "hedge Chevron lacks: about 20% of its 2025 output came from Qatar and the UAE [4], and those barrels "
     "return as prices fall in a peace scenario. This is a relative call, not a claim that either stock is cheap.")

h1("2 Background")
para("Brent was $71 just before the conflict began in late February, passed $100 by mid-March [20], averaged "
     "about $104 in the second quarter [14] and $114 in September, and was near $105 on October 13 [13][21]. "
     "The EIA's October outlook sees $105 in late 2026, $84 on average in 2027 and $74 by late 2027 [13]. Both "
     "stocks are up roughly 30% or more this year [19]. Chevron closed its $53 billion Hess acquisition in July "
     "2025, adding 30% of Guyana's Stabroek block [18]. Exxon's Middle East volumes cut its upstream earnings "
     "by $1.28 billion in the first half, about $1.03 billion of it in Q2 [5]; Chevron's regional exposure is "
     "about 1% of output [17]. Table 1 lists the inputs used throughout.")
para("This matters for a buyer today because the share prices already reflect much of the windfall. The "
     "useful question is not which company is better run, but which one still earns its share price once $105 "
     "oil becomes $74 oil. Both report third quarter results at the end of October, which gives a near-term "
     "check on the analysis.")
table("Table 1", "Key inputs. CFO is operating cash flow excluding working capital. Sources in brackets.",
      ["Item", "XOM", "CVX"],
      [["Share price, Oct 13, 2026 [16]", "$173.24", "$217.73"],
       ["Market value (Q2 diluted shares) [1][6]", f"${co['XOM']['mcap']:.0f}B", f"${co['CVX']['mcap']:.0f}B"],
       ["Dividend, annual / yield [1][6]", f"${co['XOM']['div']:.1f}B / {100*co['XOM']['div']/co['XOM']['mcap']:.1f}%",
        f"${co['CVX']['div']:.1f}B / {100*co['CVX']['div']/co['CVX']['mcap']:.1f}%"],
       ["2025 CFO [2][7]", "$59.7B", "$34.9B"],
       ["2026 capex guidance [2][9]", "$27–29B", "$18–19B"],
       ["Sensitivity per $1 Brent [3][9]", "~$0.70B", "~$0.60B"],
       ["Middle East share of output [4][17]", "~20%", "~1%"]],
      [3.1, 1.1, 1.3])

h1("3 Method")
para("Annualizing Q2 2026 cash flow would be misleading. Even after removing the oil price effect with each "
     "company's own sensitivity, Q2 implies about $19 billion (Exxon) and $23 billion (Chevron) more annual cash "
     "flow than 2025 actually produced. Much of the gap is war-inflated refining margins: Exxon puts each $1 of "
     "refining margin at about $800 million of annual earnings [5].")
para("Instead we start from 2025 cash flow (Brent $69) and move it to each scenario price using the companies' "
     "disclosed sensitivities [3][9]. For Chevron we add the part of 2025 before Hess was acquired, +$"
     f"{hess_adj:.1f} billion from Hess's 2024 results [12], using an assumed $0.10 billion per $1 Hess "
     "sensitivity (zero or double moves Chevron's breakevens by about $1). Free cash flow subtracts capex of $28 "
     "billion (Exxon) and $18.5 billion (Chevron) [9]. Buybacks are tested at the same 2.8% of market value for "
     f"both: Exxon's $20 billion plan and about ${co['CVX']['bb']:.0f} billion for Chevron. Scenarios are $60 "
     "(glut), $74 (EIA late 2027), $84 (EIA 2027 average) and $105 (war persists), where Exxon also loses $4.1 "
     "billion a year of Middle East volumes. We treat $74 as one input, not a forecast; the EIA moved its Q4 "
     "2026 view by $14 in one month [13]. The model credits neither company with growth after 2025 except Hess, "
     "which understates Chevron more.")

h1("4 Results")
figure("fig1_fcf_yield.png", "Figure 1",
       "Free cash flow yield under four Brent scenarios, with dividend yields. Author's model.")
rows = []
for B, nm in [(60, "Glut, $60"), (74, "Peace, $74"), (84, "EIA 2027, $84"), (105, "War, $105")]:
    rows.append([nm, f"${xr[B]['fcf']:.1f}B", f"${xr[B]['resid']:.1f}B", f"${cr[B]['fcf']:.1f}B", f"${cr[B]['resid']:.1f}B"])
table("Table 2", "Annual free cash flow, and cash left after dividends before buybacks. Author's model.",
      ["Scenario", "XOM FCF", "XOM after div.", "CVX FCF", "CVX after div."], rows, [1.3, 1.0, 1.1, 1.0, 1.1])
para(f"Chevron yields more at $74 ({cr[74]['fcf_yld']:.1f}% versus {xr[74]['fcf_yld']:.1f}%), $84 "
     f"({cr[84]['fcf_yld']:.1f}% versus {xr[84]['fcf_yld']:.1f}%) and $105 ({cr[105]['fcf_yld']:.1f}% versus "
     f"{xr[105]['fcf_yld']:.1f}%). Exxon wins only at $60. The two cross at about $67. The trade-off is "
     f"volatility: from war to peace, Chevron's FCF falls {100*(1-cr[74]['fcf']/cr[105]['fcf']):.0f}% and "
     f"Exxon's {100*(1-xr[74]['fcf']/xr[105]['fcf']):.0f}%. At $60, Exxon has ${xr[60]['resid']:.1f} billion "
     "left after its dividend while Chevron roughly breaks even.")
figure("fig2_breakevens.png", "Figure 2",
       "Brent price needed to fund capex and dividend, and capex, dividend and a 2.8% buyback. Author's model.")
para(f"Exxon covers capex and dividend at about ${X['be_div']:.0f} and adds its buyback at ${X['be_bb']:.0f}; "
     f"Chevron needs ${C['be_div']:.0f} and ${C['be_bb']:.0f}. Both full payout programs need more than $74, so "
     "buybacks will likely shrink after the war or lean on the balance sheet, as Exxon's did in 2025 when it "
     "paid out $37.2 billion against $26.1 billion of free cash flow [2]. Chevron says its breakeven is below $50 [10], about $10 under "
     "our estimate. If management is closer to right, Chevron's lead is larger than Figure 1 shows.")

h1("5 Alternatives and risks")
para("**The case for Exxon.** It falls less when oil drops, its dividend is safe at $60, and its Middle East "
     "barrels act as a hedge. Investors expecting oil below about $67, or who care most about dividend safety, "
     "should prefer it. **Owning neither** is also reasonable: both stocks rallied on a windfall the EIA calls "
     "temporary, and a 5% peace-case yield is not a margin of safety.")
para("**Risks to the call.** If Chevron's Q3 results imply normalized cash flow below our "
     f"${co['CVX']['base']:.1f} billion, its lead shrinks. A fast ceasefire hurts Chevron more and returns "
     "Exxon's volumes; if those volumes never return, Exxon's peace-case yield drops to about 4.5%. A post-war "
     "glut below $60 is the scenario where Exxon wins, and Chevron's Tengiz output depends on the "
     "intermittently disrupted CPC pipeline [17]. Each $1 billion error in a company's base cash flow moves its "
     "breakevens by $1.40 to $1.70 a barrel.")

h1("6 Conclusion")
para("The market is valuing both majors partly on wartime cash flow. Normalized, Chevron offers more cash and "
     "a higher dividend per dollar invested, even before credit for growth the model leaves out. Exxon is the "
     "safer stock if oil collapses. For an investor who wants the exposure and accepts the EIA's path, we "
     "prefer Chevron, and we would revisit the call after third quarter results.")

# ---------------------------------------------------------------- references
doc.add_page_break()
h1("References")
refs = [
 "ExxonMobil Holdings Corp. (2026, July 31). ExxonMobil announces second-quarter 2026 results [Form 8-K, Ex. 99.1]. https://www.sec.gov/Archives/edgar/data/0002115436/000211543626000006/livef8k2q26991.htm",
 "ExxonMobil. (2026, January 30). ExxonMobil announces 2025 results. https://corporate.exxonmobil.com/news/news-releases/2026/0130-exxonmobil-announces-2025-results",
 "ExxonMobil. (2026). 2Q 2026 sensitivities and key assumptions. https://d1io3yog0oux5.cloudfront.net/_9c0226c1dc3efd2dbe7760ece1f33886/exxonmobil/db/2431/22314/file/Sensitivities+and+Key+Assumptions.pdf",
 "Exxon Mobil Corp. (2026, April 8). Form 8-K, first-quarter 2026 earnings considerations. https://www.sec.gov/Archives/edgar/data/34088/000003408826000056/f8k1q992040826.htm",
 "ExxonMobil Holdings Corp. (2026, July 31). 2Q 2026 earnings call presentation slides (slides 9, 13, 17 and sensitivities page). https://investor.exxonmobil.com",
 "Chevron Corp. (2026, July 31). Chevron reports second quarter 2026 results [Form 8-K, Ex. 99.1]. https://www.sec.gov/Archives/edgar/data/0000093410/000009341026000162/a06302026ex9918-k.htm",
 "Chevron Corp. (2026, January). Chevron reports fourth quarter 2025 results [Form 8-K, Ex. 99.1]. https://www.sec.gov/Archives/edgar/data/93410/000009341026000019/a12312025ex9918-k.htm",
 "Chevron Corp. (2026). Form 10-Q for the quarter ended June 30, 2026. https://www.sec.gov/Archives/edgar/data/0000093410/000009341026000167/cvx-20260630.htm",
 "Chevron Corp. (2026). Form 10-Q for the quarter ended March 31, 2026. https://www.sec.gov/Archives/edgar/data/0000093410/000009341026000113/cvx-20260331.htm",
 "Chevron Corp. (2025, November 12). 2025 Investor Day, edited transcript. https://chevroncorp.gcs-web.com/static-files/f8964224-9b9a-4be6-8644-9f96bd1e1fe4",
 "Chevron Corp. (2026, August 7). Chevron 2026 investor presentation. https://chevroncorp.gcs-web.com/static-files/6fe32bf1-951a-44c8-81ee-74cd626c0095",
 "Hess Corp. (2025). Form 10-K for fiscal year 2024. https://www.sec.gov/Archives/edgar/data/4447/000162828025008716/hes-20241231.htm",
 "U.S. Energy Information Administration. (2026, October 6). Short-Term Energy Outlook, October 2026. https://www.eia.gov/outlooks/steo/archives/oct26.pdf",
 "APA. (2026). World Bank releases average oil price data for Brent and WTI. https://en.apa.az/energy-and-industry/world-bank-releases-average-oil-price-data-for-brent-and-wti-519072",
 "Statista. (2026). Brent crude oil price annually 1976–2025. https://www.statista.com/statistics/262860/weekly-crude-oil-prices",
 "Morningstar. (2026, October 13). XOM and CVX stock quotes. https://www.morningstar.com/stocks/xnys/xom/quote ; https://www.morningstar.com/stocks/xnys/cvx/quote",
 "The Motley Fool. (2026, August 3). Chevron (CVX) Q2 2026 earnings call transcript. https://www.fool.com/earnings/call-transcripts/2026/08/03/chevron-cvx-q2-2026-earnings-call-transcript/",
 "World Oil. (2025, July 18). Chevron completes $53 billion Hess takeover following lengthy Exxon arbitration. https://worldoil.com/news/2025/7/18/chevron-completes-53-billion-hess-takeover-following-lengthy-exxon-arbitration/",
 "24/7 Wall St. (2026, July 23). At $100 per barrel, which oil stock has dominated in 2026; The Motley Fool. (2026, July 30). These 3 oil stocks surged in the first half of 2026. https://247wallst.com/investing/2026/07/23/at-100-per-barrel-which-oil-stock-has-dominated-in-2026-exxonmobil-chevron-or-bp/",
 "Politis. (2026). Tracking oil prices since the beginning of the war 100 days ago. https://en.politis.com.cy/news/world/1011001/tracking-oil-prices-since-the-beginning-of-the-war-100-days-ago",
 "Trading Economics. (2026, October 13). Brent crude oil. https://tradingeconomics.com/commodity/brent-crude-oil",
]
for i, r_ in enumerate(refs, 1):
    p = doc.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.1
    rr = p.add_run(f"[{i}] {r_}"); rr.font.size = Pt(9)

doc.save("Investor_Analyst_Report.docx")
print("saved")
